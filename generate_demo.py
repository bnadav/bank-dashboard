#!/usr/bin/env python3
"""
Generate fake demo xlsx files and docs/index.html for GitHub Pages showcase.

Usage:
    python generate_demo.py

Output:
    demo_01_2026.xlsx, demo_02_2026.xlsx, demo_03_2026.xlsx
    docs/index.html
"""

import sys
import os
import random
from datetime import datetime

# sys.stdout UTF-8 wrapping is handled by bank_dashboard on import below

try:
    import openpyxl
except ImportError:
    print("openpyxl not found. Run: pip install openpyxl")
    sys.exit(1)

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bank_dashboard import (
    LOCAL_ANCHOR, FOREIGN_ANCHOR,
    find_section, parse_section, build_combined_html,
)

random.seed(42)

CARDS = ['6138', '8689']
MONTHS = [(1, 2026), (2, 2026), (3, 2026)]  # months to generate

# (merchant, charge_min, charge_max, card, txn_type, is_installment)
LOCAL_TEMPLATES = [
    ('פז תחנת דלק',        250,  400, '6138', 'עסקה רגילה',     False),
    ('Yellow פז',            40,   85, '6138', 'עסקה רגילה',     False),  # < 100 → Yellow category
    ('Yellow חניון',         25,   70, '8689', 'עסקה רגילה',     False),  # < 100 → Yellow category
    ('דרך ארץ כביש 6',       30,  120, '8689', 'עסקה רגילה',     False),
    ('מגדל ביטוח',          350,  600, '6138', 'הוראת קבע',      False),
    ('ביטוח לאומי',         200,  400, '8689', 'הוראת קבע',      False),
    ('פרטנר תקשורת',        120,  200, '6138', 'הוראת קבע',      False),
    ('HOT Mobile',           80,  150, '8689', 'הוראת קבע',      False),
    ('שופרסל דיל',          200,  500, '6138', 'עסקה רגילה',     False),
    ('רמי לוי שיווק',       150,  400, '8689', 'עסקה רגילה',     False),
    ('סופר פארם',            80,  200, '6138', 'עסקה רגילה',     False),
    ('ארומה קפה',            35,   80, '6138', 'עסקה רגילה',     False),
    ('WOLT הזמנה',           60,  150, '8689', 'עסקה רגילה',     False),
    ('רולדין',               40,  100, '6138', 'עסקה רגילה',     False),
    ('קופת חולים מכבי',      30,   80, '8689', 'עסקה רגילה',     False),
    ('בית חולים הדסה',      100,  300, '6138', 'עסקה רגילה',     False),
    ('NETFLIX',              60,   60, '6138', 'הוראת קבע',      False),
    ('GOOGLE ONE',           10,   15, '8689', 'הוראת קבע',      False),
    ('IKEA ראשון לציון',    200,  400, '6138', 'תשלומים - רגיל', True),   # installment
    ('דלתא גלריה',          100,  300, '8689', 'עסקה רגילה',     False),
    ('אולימפוס ספורט',      150,  250, '6138', 'הוראת קבע',      False),
    ('אור ירוק גן ילדים',   800, 1200, '8689', 'הוראת קבע',      False),
    ('BIT העברה',            50,  300, '6138', 'עסקה רגילה',     False),
    ('עיגול לטובה',           5,   30, '8689', 'עסקה רגילה',     False),
    ('פועלים- דמי כרטיס',   12,   12, '6138', 'דמי כרטיס',      False),
    ('פועלים- דמי כרטיס',   12,   12, '8689', 'דמי כרטיס',      False),
    ('חנות שונות בע"מ',      50,  200, '6138', 'עסקה רגילה',     False),
]

# Credit/refund transaction added per month
CREDIT_TEMPLATE = ('רמי לוי שיווק - זיכוי', -150, -40, '8689', 'זיכוי')

# (merchant, currency, ils_min, ils_max, orig_min, orig_max, card)
FOREIGN_TEMPLATES = [
    ('SPOTIFY',         'USD',  38,   42,  10.99, 10.99, '6138'),
    ('AIRBNB STAY',     'USD', 500, 1200, 140.0, 330.0,  '8689'),
    ('BOOKING.COM',     'EUR', 300,  900,  75.0, 225.0,  '6138'),
]


def rand_float(lo: float, hi: float) -> float:
    return round(random.uniform(lo, hi), 2)


def make_xlsx(path: str, month: int, year: int) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'גיליון1'

    billing_date = datetime(year, month, 28)

    # Blank header rows (like real bank exports)
    ws.append([None] * 14)
    ws.append([None] * 14)

    # ── LOCAL section ──────────────────────────────────────────────────────────
    row = [None] * 14
    row[0] = LOCAL_ANCHOR
    ws.append(row)

    row = [None] * 14
    row[0] = 'חשבון 12-782-82779'
    ws.append(row)

    ws.append(['כרטיס', 'תאריך חיוב', 'תאריך עסקה', 'שם בית עסק', 'סכום מקורי', 'סכום חיוב',
               'אסמכתא', '', '', '', '', '', 'סוג עסקה', ''])

    for merchant, lo, hi, card, txn_type, is_installment in LOCAL_TEMPLATES:
        charge = rand_float(lo, hi)
        purchase = round(charge * random.randint(3, 6), 2) if is_installment else charge
        txn_date = datetime(year, month, random.randint(1, 25))
        row = [None] * 14
        row[0]  = card
        row[1]  = billing_date
        row[2]  = txn_date
        row[3]  = merchant
        row[4]  = purchase          # col[4] = original purchase total
        row[5]  = charge            # col[5] = charge this cycle (ILS)
        row[6]  = str(random.randint(100000, 999999))
        row[12] = txn_type
        ws.append(row)

    # Credit/refund
    c_merchant, c_lo, c_hi, c_card, c_type = CREDIT_TEMPLATE
    credit_amt = rand_float(c_lo, c_hi)   # negative
    row = [None] * 14
    row[0]  = c_card
    row[1]  = billing_date
    row[2]  = datetime(year, month, random.randint(1, 25))
    row[3]  = c_merchant
    row[4]  = credit_amt
    row[5]  = credit_amt
    row[6]  = str(random.randint(100000, 999999))
    row[12] = c_type
    ws.append(row)

    # Terminator
    ws.append([None] * 14)

    # ── FOREIGN section ────────────────────────────────────────────────────────
    row = [None] * 14
    row[0] = FOREIGN_ANCHOR          # exact string — parser requires this
    ws.append(row)

    row = [None] * 14
    row[0] = 'חשבון 12-782-82779'
    ws.append(row)

    ws.append(['כרטיס', 'תאריך חיוב', 'תאריך עסקה', 'שם בית עסק', 'סכום חיוב', 'סכום מקורי',
               'מטבע', 'אסמכתא', '', '', '', '', '', 'סוג עסקה'])

    for merchant, currency, ils_lo, ils_hi, orig_lo, orig_hi, card in FOREIGN_TEMPLATES:
        charge = rand_float(ils_lo, ils_hi)
        orig   = rand_float(orig_lo, orig_hi)
        txn_date = datetime(year, month, random.randint(1, 25))
        row = [None] * 14
        row[0]  = card
        row[1]  = billing_date
        row[2]  = txn_date
        row[3]  = merchant
        row[4]  = charge            # col[4] = charge ILS (foreign layout)
        row[5]  = orig              # col[5] = original amount
        row[6]  = currency
        row[7]  = str(random.randint(100000, 999999))
        row[13] = 'עסקה רגילה'
        ws.append(row)

    # Terminator
    ws.append([None] * 14)

    wb.save(path)
    print(f'  ✓ Written: {os.path.basename(path)}')


def parse_demo_file(path: str) -> list:
    wb = openpyxl.load_workbook(path)
    rows = list(wb['גיליון1'].iter_rows(values_only=True))
    txns = []
    local_i = find_section(rows, LOCAL_ANCHOR)
    if local_i != -1:
        txns += parse_section(rows, local_i, '????', 'local')
    foreign_i = find_section(rows, FOREIGN_ANCHOR)
    if foreign_i != -1:
        txns += parse_section(rows, foreign_i, '????', 'foreign')
    return txns


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    docs_dir = os.path.join(script_dir, 'docs')
    os.makedirs(docs_dir, exist_ok=True)

    months_data = []
    for month, year in sorted(MONTHS, reverse=True):  # newest-first for tabs
        fname = f'demo_{month:02d}_{year}.xlsx'
        fpath = os.path.join(script_dir, fname)
        print(f'Generating {fname}...')
        make_xlsx(fpath, month, year)
        txns = parse_demo_file(fpath)
        charges = sum(1 for t in txns if t['charge'] > 0)
        credits = sum(1 for t in txns if t['charge'] < 0)
        print(f'  → {len(txns)} transactions ({charges} charges, {credits} credits)')
        months_data.append((txns, month, year))

    html = build_combined_html(months_data)

    # Inject demo banner into the page (not part of bank_dashboard.py)
    demo_banner = (
        '<div style="background:#FFB703;color:#1a1a2e;text-align:center;'
        'padding:10px 16px;font-size:.88rem;font-weight:600;letter-spacing:.3px;">'
        '⚠ זהו דשבורד הדגמה עם נתונים מדומים בלבד — לא נתונים אמיתיים'
        '</div>'
    )
    html = html.replace('<div class="page">', demo_banner + '\n<div class="page">', 1)

    out_path = os.path.join(docs_dir, 'index.html')
    with open(out_path, 'w', encoding='utf-8') as fh:
        fh.write(html)

    print(f'\n✓ Demo dashboard saved: docs/index.html  ({len(months_data)} months)')


if __name__ == '__main__':
    main()
